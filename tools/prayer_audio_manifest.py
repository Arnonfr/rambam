"""Write the v2 cut manifest, clip records and coverage workbook from the alignment and cuts.

Inputs: מיפוי/sections.json, מיפוי/יישור/align/*.json, מיפוי/יישור/cuts-v2.json and, when present,
מיפוי/review-decisions.json (decisions exported by a person from בדיקת-אודיו.html).
Only that file can mark a clip as approved by listening / nusach; nothing here infers approval.
v1 files are copied to מיפוי/v1 before being replaced.
"""
import json
import shutil
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from prayer_audio_pipeline import ROOT

MAP = ROOT / 'מיפוי'
WORK = MAP / 'יישור'
WORKBOOK = MAP / 'בדיקת-כיסוי-מקטעי-תפילה.xlsx'
TECH = 'נבדק טכנית'
PENDING = 'ממתין להאזנה ולאישור נוסח'

SOURCE_TITLES = {
    '639ebb9c': 'פתיחת שחרית וקרבנות', '96a701fc': 'פסוקי דזמרה — הודו עד ישתבח',
    '0ddc0de1': 'והוא רחום — תוספת שני וחמישי', '35e24669': 'סיום תחנון, אשרי, למנצח, ובא לציון, תפילה לדוד, בית יעקב, שיר המעלות',
    '9ea66621': 'הושיענו, אין כאלהינו, פיטום הקטורת, תנא דבי אליהו, עלינו, אל תירא',
    'ce85201d': 'שיר של יום — ראשון', '2fa72994': 'שיר של יום — שני', 'a75d860b': 'שיר של יום — שלישי',
    'dd64d422': 'שיר של יום — רביעי', '55eecb96': 'שיר של יום — חמישי', 'b260d7d9': 'שיר של יום — שישי',
}
EXTRA_SECTIONS = {
    'extra-hareni': 'הריני מקבל', 'extra-adon-olam': 'אדון עולם', 'extra-psalm-124': 'שיר המעלות לולי ה׳ (תהלים קכד)',
    'extra-hoshienu': 'הושיענו (סיום שיר של יום)', 'extra-ein-kelokeinu': 'אין כאלהינו',
    'extra-ketoret-end': 'פיטום הקטורת בסיום התפילה', 'extra-tana': 'תנא דבי אליהו ותלמידי חכמים',
}
# Recording-level findings for the nusach review, stated as evidence, not as conclusions.
FINDINGS = [
    ('639ebb9c', 37.8, 54.6, 'korbanot-akedah-v2', 'בהקלטה לפני העקדה: ״אלהינו ואלהי אבותינו זכרנו בזכרון טוב…״. אינו בטקסט תורה אור שבאפליקציה. התאמה אקוסטית סבירה לנוסח המקובל.'),
    ('639ebb9c', 156.9, 254.0, 'korbanot-akedah-v2', 'אחרי פסוקי העקדה כ־97 שניות טקסט נוסף. פתיחתו מתאימה ל״רבונו של עולם כמו שכבש״, אבל המשך ארוך מנוסח זה — לא זוהה במלואו. כלול בסוף מקטע העקדה.'),
    ('639ebb9c', 436.0, 451.3, 'korbanot-tamid-v2', 'בין תרומת הדשן לפרשת התמיד: ״יהי רצון… שתרחם עלינו… קרבן התמיד״. אינו בטקסט תורה אור שבאפליקציה.'),
    ('639ebb9c', 631.7, 647.9, 'korbanot-ketoret-v2', '״ה׳ צבאות עמנו״ / ״אשרי אדם״ / ״ה׳ הושיעה״ (שלוש פעמים כל אחד לפי הסידור) — התאמה אקוסטית חלשה; לבדוק כמה פעמים נאמר ובאיזה נוסח.'),
    ('639ebb9c', 689.4, 716.1, 'korbanot-ketoret-v2', 'אחרי ״אביי הוה מסדר״: ״אנא בכח״. אינו בטקסט תורה אור שבאפליקציה; כלול בסוף מקטע הקטורת.'),
    ('639ebb9c', 323.2, 343.0, 'korbanot-shema-v2', '״ברוך שם כבוד מלכותו״ לא נשמע — כמקובל נאמר בלחש. לא חסר בהקלטה.'),
    ('96a701fc', 230.9, 243.0, 'pzd-after-hodu-v2', '״ה׳ מלך ה׳ מלך״ — נמצא, אך במילים מתוחות; לבדוק בהאזנה שנאמר פעמיים כבסידור.'),
    ('96a701fc', 592.0, 598.3, 'pzd-psalm147-v2', 'סוף תהלים קמז (״לא עשה כן לכל גוי״) — התאמה חלשה בגבול עם תהלים קמח.'),
    ('9ea66621', 19.4, 22.7, None, 'אחרי ״הושיענו״: פסוק ״קוה אל ה׳…״ (תהלים כז, יד) בלבד — לא ״לדוד ה׳ אורי״ שלם. לא יוצא כמקטע.'),
    ('9ea66621', 22.7, 31.4, 'ending-ein-kelokeinu-v2', 'כ־8 שניות לפני ״אין כאלהינו״ שלא זוהו (המילה הראשונה מתוחה). ייתכן חזרה או טקסט נוסף — בדיקת גבול בהאזנה.'),
    ('9ea66621', 86.7, 102.1, 'ending-ketoret-v2', 'בסיום: ״רבן שמעון בן גמליאל אומר: הצרי…״ נמצא; ״תניא רבי נתן״ ו״בר קפרא״ לא נאמרו כאן.'),
]


def load(path):
    return json.loads(path.read_text())


def backup_v1():
    target = MAP / 'v1'
    target.mkdir(exist_ok=True)
    for name in ('cut-manifest.json', 'clips.json', 'playlist-templates.json', WORKBOOK.name):
        source = MAP / name
        if source.exists() and not (target / name).exists():
            shutil.copy2(source, target / name)


def unidentified_regions(source_id, alignment):
    """Stretches with speech (CTC character spikes) between aligned paragraphs."""
    em = np.load(WORK / 'emissions' / f'audio-{source_id}.npy')
    spikes = np.argmax(em, 1) != 0
    units = [u for u in alignment['units'] if not u.get('skipped')]
    edges = [(0.0, 0.0)] + [(u['start'], u['end']) for u in units] + [(len(em) * 0.02, len(em) * 0.02)]
    regions = []
    for (_, a), (b, _) in zip(edges, edges[1:]):
        chars = int(spikes[int(a / 0.02):int(b / 0.02)].sum())
        if b - a >= 1.5 and chars >= 10:
            regions.append(dict(source_id=source_id, start=round(a, 2), end=round(b, 2), character_spikes=chars))
    return regions


def decisions():
    path = MAP / 'review-decisions.json'
    if not path.exists():
        return {}
    return {d['id']: d for d in load(path).get('decisions', [])}


def clip_records(cuts, review):
    out = []
    for c in cuts:
        d = review.get(c['id'], {})
        listening = d.get('listening', 'pending')
        nusach = d.get('nusach', 'pending')
        approved = listening == 'approved' and nusach == 'approved' and d.get('sha256') == c['sha256']
        technical = c['decoded_ok'] and c['true_peak_dbfs'] <= -0.5 and abs(c['loudness_lufs'] + 20) <= 1.0
        out.append(dict(
            c, verification=dict(
                technical='passed' if technical else 'failed',
                text_alignment='passed_with_flags' if (c['boundary_flags'] or c['possible_unmatched_speech']) else 'passed',
                listening=listening, nusach=nusach, approved=approved,
                reviewer_notes=d.get('notes', '')),
            status=('אושר בהאזנה ובנוסח' if approved else f'{TECH} · יושר אוטומטית לטקסט · {PENDING}'),
        ))
    return out


def coverage(sections, clips):
    by_section = {}
    for c in clips:
        by_section.setdefault(c['section_id'], []).append(c)
    rows = []
    for s in sections:
        found = by_section.get(s['id'], [])
        if found:
            status = 'קיים — ' + ('אושר' if all(c['verification']['approved'] for c in found) else f'{TECH}, {PENDING}')
            if s['id'] == 'prayer-50e97cbac988':
                status += ' (א׳–ו׳; שבת מחוץ לתחום)'
        elif s['role'] == 'מניין':
            status = 'חסר — אין הקלטה (מקטע מניין; המקליט התפלל ביחיד)'
        elif s['service'] in ('בוקר', 'שחרית'):
            status = 'חסר — אין הקלטה במקורות שסופקו'
        else:
            status = 'חסר — אין הקלטות לתפילה זו'
        rows.append(dict(s, status=status, clip_ids=[c['id'] for c in found]))
    return rows


def style(sheet, widths):
    sheet.sheet_view.rightToLeft = True
    head = PatternFill('solid', fgColor='A5E1ED')
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = head
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def workbook(sources, clips, rows, regions, superseded):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'הקלטות מקור'
    ws.append(['מזהה', 'תוכן מאומת ביישור', 'משך דקות', 'מקטעים שיוצאו', 'קטעים שלא זוהו', 'קובץ מקור', 'Drive'])
    for s in sources:
        ws.append([s['id'], s['content'], round(s['duration_seconds'] / 60, 2), ', '.join(s['clip_ids']),
                   '; '.join(f'{r["start"]:.1f}–{r["end"]:.1f}' for r in s['unidentified']) or '—', s['filename'], s['drive_url']])
    style(ws, {'A': 11, 'B': 48, 'C': 9, 'D': 50, 'E': 24, 'F': 20, 'G': 30})

    ws = wb.create_sheet('מקטעי אודיו')
    ws.append(['מזהה קטע', 'תוכן', 'מזהה מקטע', 'מקור', 'התחלה במקור (ש׳)', 'סוף במקור (ש׳)', 'משך (ש׳)', 'מילים ראשונות',
               'מילים אחרונות', 'בדיקה טכנית', 'יישור לטקסט', 'האזנה', 'נוסח', 'מצב', 'נשימות שהונמכו', 'LUFS', 'True peak',
               'הערות גבול', 'דיבור שלא הותאם לטקסט', 'קובץ', 'SHA256'])
    for c in clips:
        v = c['verification']
        ws.append([c['id'], c['title'], c['section_id'], c['source_id'], c['source_start_seconds'], c['source_end_seconds'],
                   c['duration_seconds'], c['first_words'], c['last_words'], v['technical'], v['text_alignment'], v['listening'],
                   v['nusach'], c['status'], len(c['breaths_attenuated']), c['loudness_lufs'], c['true_peak_dbfs'],
                   '\n'.join(c['boundary_flags']) or '—',
                   '\n'.join(f'{u["word"]} @{u["source_start"]:.1f} ({u["seconds"]}ש׳)' for u in c['possible_unmatched_speech']) or '—',
                   c['file'], c['sha256']])
        ws.cell(ws.max_row, 20).hyperlink = '../' + c['file']
    style(ws, {'A': 26, 'B': 22, 'C': 21, 'D': 10, 'H': 26, 'I': 26, 'N': 34, 'R': 36, 'S': 30, 'T': 30})

    ws = wb.create_sheet('כיסוי מול האקסל')
    ws.append(['מס׳ באקסל', 'מזהה קבוע', 'תפילה', 'קבוצה', 'מקטע', 'מתי', 'התאמות', 'תפקיד', 'מצב כיסוי', 'קבצי אודיו'])
    for r in rows:
        ws.append([r['excel_row'], r['id'], r['service'], r['group'], r['title'], r['when'], r['exceptions'], r['role'],
                   r['status'], ', '.join(r['clip_ids']) or '—'])
    style(ws, {'B': 21, 'E': 24, 'F': 24, 'G': 30, 'I': 40, 'J': 40})

    ws = wb.create_sheet('חסרים')
    ws.append(['מס׳ באקסל', 'תפילה', 'מקטע', 'מתי', 'תפקיד', 'מצב', 'השפעה על נגינה רציפה'])
    for r in rows:
        if r['clip_ids']:
            continue
        impact = '—'
        if r['service'] in ('בוקר', 'שחרית') and r['role'] != 'מניין':
            impact = 'הנגן עוצר לפני מקטע זה ומודיע; אין דילוג שקט'
        elif r['role'] == 'מניין':
            impact = 'במצב יחיד אינו חלק מהרצף; במצב מניין הנגן עוצר כאן'
        ws.append([r['excel_row'], r['service'], r['title'], r['when'], r['role'], r['status'], impact])
    style(ws, {'C': 26, 'D': 30, 'F': 40, 'G': 44})

    ws = wb.create_sheet('מקטעים מעבר לאקסל')
    ws.append(['שם', 'מזהה מקטע', 'קובץ', 'מצב', 'הערה'])
    for section, title in EXTRA_SECTIONS.items():
        found = [c for c in clips if c['section_id'] == section]
        ws.append([title, section, ', '.join(c['id'] for c in found) or '—', found[0]['status'] if found else 'חסר',
                   'מקטע בהקלטה שאין לו שורה נפרדת באקסל המקורי'])
    style(ws, {'A': 30, 'B': 22, 'C': 28, 'D': 40, 'E': 40})

    ws = wb.create_sheet('ממצאי נוסח')
    ws.append(['מקור', 'מ־(ש׳)', 'עד (ש׳)', 'קטע', 'ממצא — לבדיקה בהאזנה מול תהלת ה׳'])
    for source_id, a, b, clip, text in FINDINGS:
        ws.append([source_id, a, b, clip or 'לא יוצא', text])
    for r in regions:
        ws.append([r['source_id'], r['start'], r['end'], '—', f'דיבור שלא הותאם לאף פסקה בטקסט ההשוואה ({r["character_spikes"]} סימני דיבור)'])
    style(ws, {'A': 10, 'D': 26, 'E': 90})

    ws = wb.create_sheet('חלופות והשלמות')
    ws.append(['מקטע', 'חלופה', 'מצב', 'קובץ', 'הערה'])
    for c in clips:
        if c.get('weekday'):
            ws.append(['שיר של יום', c['title'].split('— ')[-1], c['status'], c['id'], 'כולל פתיחת ״היום יום…״ והמזמור המלא'])
    for row in [('שיר של יום', 'שבת', 'מחוץ לתחום', '—', 'שבתות אינן בתחום השימוש'),
                ('ברכי נפשי', 'ראש חודש', 'חסר', '—', 'נדרש אחרי שיר של יום בראש חודש'),
                ('לדוד ה׳ אורי', 'אלול עד הושענא רבה', 'חסר', '—', 'קיים רק הפסוק האחרון בהקלטה 9ea66621 (19.4–22.7); לא מקטע שלם'),
                ('הלל', 'ראש חודש / חנוכה / חול המועד', 'חסר', '—', ''),
                ('מוסף', 'ראש חודש / חול המועד', 'חסר', '—', ''),
                ('עמידה', 'קיץ / חורף / תוספות היום', 'חסר', '—', 'אין הקלטת עמידה כלל')]:
        ws.append(list(row))
    style(ws, {'A': 16, 'B': 26, 'C': 40, 'D': 18, 'E': 50})

    ws = wb.create_sheet('גרסאות קודמות')
    ws.append(['מזהה v1', 'הוחלף ב־', 'הערה'])
    for old, new in superseded:
        ws.append([old, new or '—', 'חיתוך טיוטה ראשון; הקובץ נשמר בתיקיית מקטעים ולא נמחק'])
    style(ws, {'A': 32, 'B': 32, 'C': 60})
    wb.save(WORKBOOK)


SUPERSEDED = [('shir-sun-v1', 'shir-sun-v2'), ('shir-mon-v1', 'shir-mon-v2'), ('shir-tue-v1', 'shir-tue-v2'),
              ('shir-wed-v1', 'shir-wed-v2'), ('shir-thu-v1', 'shir-thu-v2'), ('shir-fri-v1', 'shir-fri-v2'),
              ('tachanun-mon-thu-v1', 'tachanun-mon-thu-v2'), ('closing-short-tachanun-v1', 'tachanun-end-v2'),
              ('closing-ashrei-v1', 'closing-ashrei-v2'), ('closing-lamnatzeach-v1', 'closing-lamnatzeach-v2'),
              ('closing-uvaletzion-v1', 'closing-uvaletzion-v2'), ('closing-tefilah-ledavid-v1', 'closing-tefilah-ledavid-v2'),
              ('closing-beit-yaakov-combined-v1', 'closing-beit-yaakov-v2 + closing-psalm124-v2'),
              ('ending-before-ketoret-v1', 'closing-hoshienu-v2'), ('ending-ein-kelokeinu-v1', 'ending-ein-kelokeinu-v2'),
              ('ending-ketoret-v1', 'ending-ketoret-v2'), ('ending-tana-v1', 'ending-tana-v2'),
              ('ending-aleinu-v1', 'ending-aleinu-v2'), ('ending-final-verses-v1', 'ending-final-verses-v2'),
              ('morning-hareni-v1', 'morning-hareni-v2'), ('morning-ma-tovu-v1', 'morning-ma-tovu-v2')]


def main():
    backup_v1()
    sections = load(MAP / 'sections.json')['sections']
    inventory = {s['id']: s for s in load(MAP / 'source-inventory.json')}
    clips = clip_records(load(WORK / 'cuts-v2.json'), decisions())
    regions, sources = [], []
    for source_id, title in SOURCE_TITLES.items():
        alignment = load(WORK / 'align' / f'{source_id}.json')
        found = unidentified_regions(source_id, alignment)
        regions += found
        sources.append(dict(id=source_id, content=title, duration_seconds=inventory[source_id]['duration_seconds'],
                            filename=inventory[source_id]['filename'], drive_url=inventory[source_id]['drive_url'],
                            sha256=inventory[source_id]['sha256'], unidentified=found,
                            units_found=[u['id'] for u in alignment['units'] if not u.get('skipped')],
                            units_not_found=[u['id'] for u in alignment['units'] if u.get('skipped')],
                            clip_ids=[c['id'] for c in clips if c['source_id'] == source_id]))
    rows = coverage(sections, clips)
    manifest = dict(
        schema_version=2,
        status='v2_cuts_technically_checked_listening_and_nusach_pending',
        method=('MMS CTC forced alignment of each source against the Torah Or text (and documented comparison '
                'texts), cut in the silence after each section from the original audio (clear microphone, no noise reduction or compression), breaths attenuated 12 dB, static gain to -20 LUFS with a peak-only limiter'),
        comparison_text='סידור תורה אור (ויקיטקסט) — אינו מהדורת תהלת ה׳; אימות נוסח מול תהלת ה׳ טרם בוצע',
        approval_rule='approved=true only from מיפוי/review-decisions.json with listening and nusach approved for the same SHA256',
        sources=sources, clips=clips, findings=[dict(source_id=s, start=a, end=b, clip=c, note=t) for s, a, b, c, t in FINDINGS],
        unidentified_regions=regions, superseded=[dict(v1=a, v2=b) for a, b in SUPERSEDED],
        coverage=[dict(section_id=r['id'], excel_row=r['excel_row'], title=r['title'], status=r['status'], clip_ids=r['clip_ids'])
                  for r in rows])
    (MAP / 'cut-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=1))
    (MAP / 'clips.json').write_text(json.dumps(dict(schema_version=2, publication_status='not_approved', clips=[
        {k: c[k] for k in ('id', 'title', 'section_id', 'source_id', 'source_start_seconds', 'source_end_seconds',
                           'duration_seconds', 'file', 'bytes', 'sha256', 'decoded_ok', 'loudness_lufs', 'true_peak_dbfs',
                           'verification', 'status') + (('weekday',) if 'weekday' in c else ())}
        for c in clips]), ensure_ascii=False, indent=1))
    workbook(sources, clips, rows, regions, SUPERSEDED)
    present = sum(1 for r in rows if r['clip_ids'])
    print(f'{len(clips)} clips; {present}/{len(rows)} sheet rows have audio; '
          f'{sum(1 for c in clips if c["verification"]["approved"])} approved; {len(regions)} unidentified regions')


if __name__ == '__main__':
    main()
