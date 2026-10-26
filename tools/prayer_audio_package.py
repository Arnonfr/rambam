"""Export explicit, reviewable cuts and workbook coverage without altering originals."""
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from prayer_audio_pipeline import ROOT, ffmpeg, save

PLAN = ROOT.parent / 'מסמכים' / 'מיפוי-מקטעי-תפילה.xlsx'
MANIFEST = ROOT / 'מיפוי' / 'cut-manifest.json'


def sections():
    wb = openpyxl.load_workbook(PLAN, data_only=True)
    out = []
    for row in wb.worksheets[0].iter_rows(min_row=6, values_only=True):
        if not isinstance(row[0], int):
            continue
        stable = hashlib.sha256('|'.join(str(x) for x in row[1:4]).encode()).hexdigest()[:12]
        out.append(dict(id=f'prayer-{stable}', excel_row=row[0], service=row[1], group=row[2],
                        title=row[3], when=row[4], exceptions=row[5], role=row[6]))
    return out


def export():
    inv = {x['id']: x for x in json.loads((ROOT / 'מיפוי/source-inventory.json').read_text())}
    manifest = json.loads(MANIFEST.read_text())
    if len({x['id'] for x in manifest['clips']}) != len(manifest['clips']):
        raise ValueError('Duplicate clip ID')
    previous_path = ROOT / 'מיפוי' / 'clips.json'
    previous = {x['id']: x for x in json.loads(previous_path.read_text())['clips']} if previous_path.exists() else {}
    target_dir = ROOT / 'מקטעים'
    target_dir.mkdir(exist_ok=True)
    records = []
    for clip in manifest['clips']:
        source = inv[clip['source_id']]
        start, end = float(clip['start_seconds']), float(clip['end_seconds'])
        if not 0 <= start < end <= source['duration_seconds'] + 0.15:
            raise ValueError(f'Invalid source range: {clip["id"]}')
        target = target_dir / (clip['id'] + '.m4a')
        master = ROOT / 'עיבוד' / ('audio-' + clip['source_id'] + '-clean.flac')
        if target.exists():
            old = previous.get(clip['id'])
            if not old or any(old.get(k) != clip[k] for k in ['source_id', 'start_seconds', 'end_seconds']):
                raise ValueError(f'Existing export has unverified provenance: {clip["id"]}; use a new versioned ID')
        # Decode / re-encode for accurate boundaries; no packet-copy AAC seek.
        if not target.exists():
            ffmpeg('-n', '-ss', start, '-i', master, '-t', end-start, '-ac', '1',
                   '-ar', '48000', '-c:a', 'aac', '-b:a', '80k', '-movflags', '+faststart', target)
        # Full decoder validation, not just file existence.
        ffmpeg('-v', 'error', '-i', target, '-f', 'null', '-')
        metadata = subprocess.check_output(['rtk', 'proxy', 'afinfo', str(target)], text=True)
        duration = float(re.search(r'estimated duration: ([0-9.]+)', metadata).group(1))
        if abs(duration - (end-start)) > 0.2:
            raise ValueError(f'Unexpected export duration: {clip["id"]}: {duration}')
        records.append(dict(**clip, file=str(target.relative_to(ROOT)), duration_seconds=duration,
                            bytes=target.stat().st_size, sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                            source_sha256=source['sha256'], decoded_ok=True))
    save(ROOT / 'מיפוי' / 'clips.json', dict(schema_version=1, publication_status='draft', clips=records))
    print(f'Validated {len(records)} exports')


def report():
    manifest = json.loads(MANIFEST.read_text())
    clips_file = ROOT / 'מיפוי' / 'clips.json'
    clips = json.loads(clips_file.read_text())['clips'] if clips_file.exists() else manifest['clips']
    source_info = manifest.get('sources', {})
    inv = json.loads((ROOT / 'מיפוי/source-inventory.json').read_text())
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243E49')
            cell.alignment = Alignment(horizontal='right', wrap_text=True)
        ws.row_dimensions[1].height = 34
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='F3F7F8' if cell.row % 2 == 0 else 'FFFFFF')
            ws.row_dimensions[row[0].row].height = 48
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        return ws

    sheet('הקלטות מקור', ['מזהה', 'תוכן שזוהה', 'משך דקות', 'רמת זיהוי', 'הערה', 'קובץ מקור', 'Drive'], [
        [x['id'], source_info.get(x['id'], {}).get('title', 'טרם זוהה'), round(x['duration_seconds']/60, 2),
         source_info.get(x['id'], {}).get('status', 'טרם נבדק'), source_info.get(x['id'], {}).get('notes', ''),
         x['filename'], x['drive_url']] for x in inv], [16, 45, 14, 24, 65, 30, 60])
    clip_sheet = sheet('מקטעי אודיו', ['מזהה קטע', 'תוכן', 'מקור', 'התחלה שניות', 'סוף שניות', 'משך שניות', 'קובץ', 'אימות', 'הערה'], [
        [x['id'], x['title'], x['source_id'], x['start_seconds'], x['end_seconds'],
         round(x.get('duration_seconds', x['end_seconds']-x['start_seconds']), 2), x.get('file', ''),
         x.get('review_status', 'טיוטה — נדרשת האזנה'), x.get('notes', '')] for x in clips],
         [32, 40, 16, 18, 18, 18, 50, 35, 60])
    for row_index, item in enumerate(clips, 2):
        if item.get('file'):
            cell = clip_sheet.cell(row_index, 7)
            cell.hyperlink = (ROOT / item['file']).as_uri()
            cell.font = Font(name='Arial', color='0563C1', underline='single')
    sec = sections()
    covered_rows = []
    for s in sec:
        found = [x for x in clips if s['excel_row'] in x.get('excel_rows', [])]
        probable_sources = [k for k,v in source_info.items() if s['excel_row'] in v.get('excel_rows_probable', [])]
        status = 'קיים בטיוטה — אימות נדרש' if found else manifest.get('coverage_default', 'טרם נמצא במיפוי')
        if probable_sources and not found:
            status = 'זוהה בתוך רצף מקור — עדיין אין חיתוך נפרד'
        if any(x.get('coverage') == 'combined_requires_split' for x in found):
            status = 'קיים רק בתוך קובץ משולב — פיצול נדרש'
        if s['excel_row'] == 65 and found:
            status = 'כיסוי חלקי בטיוטה — פסוקי סיום נוספים טרם מופו'
        if s['title'] == 'שיר של יום':
            status = 'כיסוי לפי ימות השבוע — ראו גיליון חלופות'
        covered_rows.append([s['excel_row'], s['id'], s['service'], s['group'], s['title'], s['when'],
                             s['exceptions'], s['role'], status, ', '.join(x['id'] for x in found), ', '.join(probable_sources)])
    sheet('כיסוי מול האקסל', ['מס׳ באקסל', 'מזהה קבוע', 'תפילה', 'קבוצה', 'מקטע', 'מתי', 'התאמות', 'תפקיד', 'מצב כיסוי', 'קבצי אודיו', 'מקור שממתין לפיצול'],
          covered_rows, [14, 26, 20, 25, 35, 50, 65, 22, 45, 65, 35])
    sheet('מקטעים מעבר לאקסל', ['שם', 'מזהה קטע', 'מצב', 'הערה'], [
        [x['title'], x['id'], x.get('review_status','טיוטה'), x.get('notes','נדרש להוסיף למפת התוכן לאחר אימות')]
        for x in clips if not x.get('excel_rows')], [45, 35, 45, 70])
    sheet('חלופות והשלמות', ['מקטע', 'חלופה', 'מצב', 'קובץ', 'הערה'], manifest.get('variants', []), [32, 32, 40, 40, 70])
    quality = []
    for path in sorted((ROOT / 'מיפוי/איכות').glob('*.json')):
        item = json.loads(path.read_text())
        quality.append([item['source'], item['before']['input_i'], item['after']['output_i'],
                        item['after']['output_tp'], item['filters'], 'טרם נבדק בהאזנה', item['master']])
    sheet('איכות', ['מקור', 'לפני יישור עוצמה LUFS', 'אחרי יישור עוצמה LUFS', 'True Peak אחרי dBTP', 'סינון', 'בדיקת האזנה', 'עותק עיבוד'],
          quality, [32, 22, 22, 25, 52, 32, 40])
    wb.save(ROOT / 'מיפוי/בדיקת-כיסוי-מקטעי-תפילה.xlsx')
    save(ROOT / 'מיפוי/sections.json', dict(schema_version=1, sections=sec))
    print(f'Coverage workbook: {len(sec)} planned sections, {len(clips)} draft clips')


if __name__ == '__main__':
    {'export': export, 'report': report}[sys.argv[1]]()
