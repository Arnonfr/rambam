"""Limit the planning workbook to days when the app is used, preserving styles."""
from pathlib import Path
from copy import copy
from openpyxl import load_workbook

path = Path('/Users/hyh/Documents/Codex/לימוד-יומי/מסמכים/מיפוי-מקטעי-תפילה.xlsx')
book = load_workbook(path)
scope = 'תחום שימוש: ימי חול וחול המועד בלבד; שבתות וימים טובים מחוץ למיפוי. ערב ומוצאי שבת/חג נשארים כשהם ימי חול.'

def remove_rows(sheet, predicate):
    removed = []
    for row in range(sheet.max_row, 5, -1):
        values = [cell.value for cell in sheet[row]]
        if predicate(values):
            removed.append(values)
            sheet.delete_rows(row)
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f'A5:{sheet.cell(5, sheet.max_column).column_letter}{sheet.max_row}'
    return removed

removed_segments = remove_rows(book['מקטעי תפילה'], lambda v:
    v[1] in {'שבת', 'מחזור'} or v[2] == 'תוספות שבת וחג' or v[3] == 'קידוש' or
    v[3] in {'מוסף שבת', 'מוסף שבת וראש חודש', 'מוסף יום טוב'})
remove_rows(book['מטריצת ימים'], lambda v:
    str(v[0]).startswith('שבת') or v[0] in {'יום טוב', 'ראש השנה/יום כיפור'})
remove_rows(book['שיר של יום'], lambda v: v[0] == 'שבת')
for sheet in book:
    if not str(sheet['A3'].value or '').startswith(scope):
        sheet['A3'] = scope + ' ' + str(sheet['A3'].value or '')

sheet = book['מקטעי תפילה']
for row in range(6, sheet.max_row + 1):
    sheet.cell(row, 1, row - 5)
    name = sheet.cell(row, 4).value
    if name == 'מוסף ראש חודש':
        sheet.cell(row, 5, 'ראש חודש ביום חול בלבד')
        sheet.cell(row, 6, 'שבת ראש חודש מחוץ לתחום השימוש באפליקציה')
    elif name == 'מוסף חול המועד':
        sheet.cell(row, 5, 'חול המועד פסח או סוכות, כאשר אינו שבת')
        sheet.cell(row, 6, 'נוסח לפי המועד; שבת חול המועד מחוץ לתחום השימוש')

sheet = book['מטריצת ימים']
for row in range(6, sheet.max_row + 1):
    if str(sheet.cell(row, 1).value).startswith('חול המועד'):
        sheet.cell(row, 11, str(sheet.cell(row, 11).value or '') + '; לא בשבת; חלוקת ימי המועד לפי ישראל/חו״ל')

# Keep pre-Shabbat and pre-festival effects; these apply on permitted app days.
sheet = book['תחנון וחריגים']
for row in range(6, sheet.max_row + 1):
    if sheet.cell(row, 2).value == 'שבת':
        sheet.cell(row, 2, 'ערב שבת — מנחה')
        sheet.cell(row, 3, 'יום שישי לפני כניסת שבת')
        sheet.cell(row, 5, 'שבת עצמה אינה במיפוי; נשמר כלל מנחת ערב שבת')
    elif sheet.cell(row, 2).value == 'ראש השנה':
        sheet.cell(row, 2, 'ערב ראש השנה — השפעה על תחנון')
        sheet.cell(row, 3, 'כ״ט אלול ביום חול; א׳–ב׳ תשרי מחוץ לתחום השימוש')
        sheet.cell(row, 5, 'ראש השנה עצמו אינו במיפוי; ערב ר״ה נשאר')

book.save(path)
check = load_workbook(path)
assert not any(str(row[0]).startswith('שבת') for row in list(check['מטריצת ימים'].values)[5:])
assert sum(str(row[0]).startswith('חול המועד') for row in check['מטריצת ימים'].values) == 2
assert any(row[3] == 'מוסף חול המועד' for row in check['מקטעי תפילה'].values)
assert not any(row[0] == 'שבת' for row in check['שיר של יום'].values)
print(f'Updated {path}; removed {len(removed_segments)} Shabbat/Yom Tov-only segments; retained both Chol Hamoed profiles.')
