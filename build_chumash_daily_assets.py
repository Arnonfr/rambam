"""Build all daily portions from the supplied Chabad workbook, by verse key.

The workbook supplies boundaries only. Torah and Rashi remain Sefaria text;
previously bundled verses are preserved exactly. Downloads are cached so
rebuilding does not require a network connection.
"""
import concurrent.futures
import copy
import json
import re
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from build_chumash_rashi_assets import clean

ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT.parent / "מסמכים/chabad_chumash_daily.xlsx"
ASSET = ROOT / "app/src/main/assets/chumash_cache.json"
CACHE = ROOT.parent / "מחקר/chumash-daily/sources"
BOOK_ENDS = {"Genesis": (50, 26), "Exodus": (40, 38),
             "Leviticus": (27, 34), "Numbers": (36, 13), "Deuteronomy": (34, 12)}
DAY_NAMES = ["יום ראשון", "יום שני", "יום שלישי", "יום רביעי", "יום חמישי", "יום שישי", "שבת קודש"]


def read_catalog():
    workbook = load_workbook(WORKBOOK, data_only=True)
    rows = []
    for name in ("פרשות יחידות", "פרשות מחוברות"):
        sheet = workbook[name]
        headers = [cell.value for cell in sheet[4]]
        for values in sheet.iter_rows(min_row=5, values_only=True):
            if values[0] is not None:
                row = dict(zip(headers, values))
                row["verified_date"] = row["verified_date"].date().isoformat()
                rows.append(row)
    assert len(rows) == 427 and len({r["segment_id"] for r in rows}) == 427
    return rows


def download(book, rashi=False):
    CACHE.mkdir(parents=True, exist_ok=True)
    name = ("Rashi_on_" if rashi else "") + book
    path = CACHE / (name + ".json")
    if path.exists():
        return json.loads(path.read_text())
    c, v = BOOK_ENDS[book]
    url = f"https://www.sefaria.org/api/texts/{name}.1.1-{c}.{v}?context=0&commentary=0"
    request = urllib.request.Request(url, headers={"User-Agent": "DailyStudyApp/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        result = json.load(response)
    assert result["sections"][:2] == [1, 1], (url, result.get("sections"))
    assert result["toSections"][:2] == [c, v], (url, result.get("toSections"))
    path.write_text(json.dumps(result, ensure_ascii=False) + "\n")
    print("Downloaded", name, flush=True)
    return result


def range_keys(book, start, end, lengths):
    sc, sv = start
    ec, ev = end
    if not (1 <= sc <= ec <= len(lengths[book])):
        raise ValueError(f"Invalid chapters {book} {start}-{end}")
    if not (1 <= sv <= lengths[book][sc - 1] and 1 <= ev <= lengths[book][ec - 1]):
        raise ValueError(f"Missing endpoint {book} {start}-{end}")
    keys = [(book, c, v) for c in range(sc, ec + 1)
            for v in range(sv if c == sc else 1, (ev if c == ec else lengths[book][c - 1]) + 1)]
    if not keys or keys[0][1:] != start or keys[-1][1:] != end:
        raise ValueError(f"Invalid range {book} {start}-{end}")
    return keys


def ref_keys(ref, lengths):
    match = re.fullmatch(r"(\w+) (\d+):(\d+)-(?:([0-9]+):)?([0-9]+)", ref)
    if not match:
        raise ValueError(f"Unrecognized reference {ref}")
    book, sc, sv, ec, ev = match.groups()
    return range_keys(book, (int(sc), int(sv)), (int(ec or sc), int(ev)), lengths)


def main():
    catalog = read_catalog()
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        tasks = {(book, rashi): pool.submit(download, book, rashi)
                 for book in BOOK_ENDS for rashi in (False, True)}
        sources = {key: job.result() for key, job in tasks.items()}
    lengths = {book: [len(c) for c in sources[book, False]["he"]] for book in BOOK_ENDS}
    verses = {}
    for book in BOOK_ENDS:
        torah = sources[book, False]["he"]
        rashi = sources[book, True]["he"]
        assert len(torah) == BOOK_ENDS[book][0] and len(torah[-1]) == BOOK_ENDS[book][1]
        assert len(rashi) == len(torah), f"{book}: missing Rashi chapters"
        for c, chapter in enumerate(torah, 1):
            # Some Sefaria chapters have trailing empty slots beyond the Torah
            # chapter. Only empty padding may be dropped; never a commentary.
            assert not any(rashi[c - 1][len(chapter):]), f"{book} {c}: extra Rashi verse"
            for v, text in enumerate(chapter, 1):
                text = clean(text)
                if not text:
                    raise ValueError(f"Missing Torah verse {book} {c}:{v}")
                comments = rashi[c - 1][v - 1] if v <= len(rashi[c - 1]) else []
                assert isinstance(comments, list), f"{book} {c}:{v}: ambiguous Rashi"
                verses[book, c, v] = {"chapterNumber": c, "verseNumber": v,
                    "textWithTeamim": text, "textPlainNikud": re.sub(r"[\u0591-\u05af]", "", text),
                    "rashi": [clean(t) for t in comments if clean(t)]}

    # Reuse the exact local Torah/Rashi text wherever it already exists.
    original = json.loads(ASSET.read_text())
    reused = set()
    for parasha in original.values():
        for aliya in parasha["aliyot"]:
            keys = ref_keys(aliya["ref"], lengths)
            if len(keys) != len(aliya["verses"]):
                raise ValueError(f"{aliya['ref']}: expected {len(keys)}, got {len(aliya['verses'])}")
            for key, verse in zip(keys, aliya["verses"]):
                content = copy.deepcopy(verse)
                content.update(chapterNumber=key[1], verseNumber=key[2])
                if key in reused and verses[key] != content:
                    raise ValueError(f"Conflicting local text {key}")
                verses[key] = content
                reused.add(key)

    groups = defaultdict(list)
    for row in catalog:
        groups[row["parasha_id"]].append(row)
    result = {}
    single_keys = []
    for pid, rows in sorted(groups.items()):
        rows.sort(key=lambda r: r["day_index"])
        assert [r["day_index"] for r in rows] == list(range(1, 8)), pid
        all_keys, aliyot = [], []
        for row in rows:
            book = row["book_id"]
            start = row["start_chapter"], row["start_verse"]
            end = row["end_chapter"], row["end_verse"]
            keys = range_keys(book, start, end, lengths)
            ref = f"{book} {start[0]}:{start[1]}-{end[0]}:{end[1]}"
            aliyot.append({"segmentId": row["segment_id"], "aliyaIndex": row["day_index"],
                "aliyaName": row["day_he"], "dayName": DAY_NAMES[row["day_index"] - 1],
                "ref": ref, "heRef": row["range_he"], "sourceUrl": row["source_url"],
                "verseKeys": [f"{b}:{c}:{v}" for b, c, v in keys],
                "verses": [verses[key] for key in keys]})
            all_keys.extend(keys)
        full_keys = range_keys(book, all_keys[0][1:], all_keys[-1][1:], lengths)
        assert all_keys == full_keys, f"{pid}: gap, overlap or incorrect order"
        if "_" not in pid:
            single_keys.extend(all_keys)
        else:
            component_keys = []
            for component in pid.split("_"):
                for row in sorted(groups[component], key=lambda r: r["day_index"]):
                    component_keys.extend(range_keys(book, (row["start_chapter"], row["start_verse"]),
                        (row["end_chapter"], row["end_verse"]), lengths))
            assert all_keys == component_keys, f"{pid}: combined portion coverage differs"
        result[rows[0]["parasha_he"]] = {"contentSchemaVersion": 3, "parashaId": pid,
            "parasha": rows[0]["parasha_he"], "book": rows[0]["book_he"],
            "fullRef": f"{book} {all_keys[0][1]}:{all_keys[0][2]}-{all_keys[-1][1]}:{all_keys[-1][2]}",
            "aliyot": aliyot}
    assert single_keys == list(verses), "Single portions do not cover the Torah exactly once"
    output = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    temp = ASSET.with_suffix(".tmp")
    temp.write_text(output)
    temp.replace(ASSET)
    (ROOT / "app/src/main/assets/chumash_daily_catalog.json").write_text(
        json.dumps({"schemaVersion": 1, "segments": catalog, "chapterLengths": lengths}, ensure_ascii=False, indent=2) + "\n")
    print(f"Validated {len(verses)} verses, {len(catalog)} portions, {len(groups)} groups; preserved {len(reused)} local verses")


if __name__ == "__main__":
    main()
